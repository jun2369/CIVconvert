"""The general reader: invoices of suppliers without a customer profile, PDF and Excel."""

import base64
import re
import io
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from civconvert import generic, preview
from civconvert.extract import PdfExtract, RowBox
from civconvert.pipeline import ConversionError, InputFile, convert
from civconvert.profiles import load_generic
from pdfgen import _document, _right, _text, invoice_pdf, offset_invoice_pdf, problem_invoice, text_invoice_pdf

XLS = Path(__file__).resolve().parent / "fixtures" / "invoice.xls"


def _dfe(result) -> list[tuple[str, str, str]]:
    return [(row[3], row[4], row[5]) for row in result.rows]


def _text_items() -> list[dict[str, str]]:
    rows = [("7208948", "GG ARTICLE II CLR REPL", "LENS [CLR]OS 26", "9001.50.0000", "60", "7.5400"),
            ("7157648", "GG ARTICLE II REPL LENS", "VIV COP 24", "9001.50.0000", "60", "11.7900"),
            ("7172096", "GG RINGO STK BK V ROY", "26 PAIRS", "9004.90.0090", "48", "13.9400"),
            ("7110191", "GG RINGO RPL LENS VIV", "JET BLK", "9001.50.0000", "60", "8.9400"),
            ("7110187", "GG RINGO RPL LENS VIV", "PNK", "9001.50.0000", "12", "11.5500"),
            ("7133153", "GG RINGO/MILLIE REPL", "LENS VIV SMK 21", "9001.50.0000", "60", "8.9400")]
    return [{"po": "4500628817", "item": item, "desc": desc, "more": more, "hts": hts, "hts_desc": "sunglasses and",
             "qty": qty, "price": price, "amount": f"{int(qty) * float(price):,.2f}"} for item, desc, more, hts, qty, price in rows]


def test_invoice_without_ruling_lines_over_two_pages():
    """A packing list first, then the invoice: no lines, wrapped descriptions, the header only on its
    first page, a description that goes on at the top of the next page, a total under every page."""
    items = _text_items()
    r = convert([InputFile("vista.pdf", text_invoice_pdf(items, per_page=4))])
    assert r.customer.id == "auto"
    assert _dfe(r) == [(i["item"], f"{i['desc']} {i['more']}", i["hts"]) for i in items]
    assert [row[0] for row in r.rows] == ["1", "2", "3", "4", "5", "6"]
    assert r.document == {"invoice_no": "INV-5531", "invoice_date": "10/13/2026"}
    assert [i["_source"][-3:] for i in r.items] == ["p.2"] * 4 + ["p.3"] * 2  # the packing list (page 1) is not read
    assert {c.name: c.ok for c in r.checks} == {"Line items": True, "Subtotal qty": True, "Subtotal amount": True, "Line math": True}
    assert r.warnings == [] and not r.needs_review
    # the preview finds every line on its page
    assert [b.page for b in r.parts[0].rows] == [2, 2, 2, 2, 3, 3]


def invoice_xlsx() -> bytes:
    wb = Workbook()
    contract = wb.active
    contract.title = "Contract"  # the same goods as a contract: no HS codes – not the invoice
    contract["B2"] = "CONTRACT"
    for c, h in enumerate(["ITEM", "Model", "PRODUCT DESCRIPTION", "QUANTITY\nPCS", "UNIT PRICE\nUSD", "NET AMOUNT\nUSD"], 2):
        contract.cell(row=5, column=c, value=h)
    for r, row in enumerate([[1, "UC200T-EM", "UC型WCDMA无线通信模块", 300, 8, 2400], [2, "M95", "GSM模块", 700, 2.5, 1750]], 6):
        for c, v in enumerate(row, 2):
            contract.cell(row=r, column=c, value=v)
    inv = wb.create_sheet("Invoice")
    inv["D3"] = "QUECTEL TEST PTE. LTD.\n6 RAFFLES QUAY SINGAPORE"
    inv["B4"] = "COMMERCIAL INVOICE"
    inv["H6"], inv["I6"] = "DATE:", datetime(2026, 10, 15)
    inv["H7"], inv["I7"] = "INVOICE NO.：", 880101773
    header = ["ITEM", "Model", "HS CODE", "Ordering Code", "PRODUCT DESCRIPTION（EN）", "QUANTITY\nPCS", "UNIT PRICE\nUSD", "NET AMOUNT\nUSD"]
    rows = [
        [1, "MLG0603P2N2BTZ10", 8504500000, "QT0014103066", "INDUCTOR", 495000, 0.0012000000000000001, 594],
        [2, "0201B271K500NT", 8532241000, "QT0012002690", "电容 CAPACITOR", 30000, 0.00039999999999999996, 11.999999999999998],
        [3, "BEEF-01", 201300000, "QT1", "FROZEN BEEF", 100, 2.5, 250],  # HS 0201.30 typed as a number
        [None, None, None, None, "(BONELESS)", None, None, None],  # wrapped description
        ["TOTAL ", None, None, None, None, 525100, None, 856],
    ]
    for c, h in enumerate(header, 2):
        inv.cell(row=18, column=c, value=h)
    for r, row in enumerate(rows, 19):
        for c, v in enumerate(row, 2):
            inv.cell(row=r, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_invoice():
    r = convert([InputFile("CO-AX invoice.xlsx", invoice_xlsx())])
    assert r.customer.id == "auto"
    assert r.filename == "CargoWise Input - CO-AX invoice.csv"
    assert _dfe(r) == [("MLG0603P2N2BTZ10", "INDUCTOR", "8504500000"),
                       ("0201B271K500NT", "CAPACITOR", "8532241000"),  # Chinese removed, as for every customer
                       ("BEEF-01", "FROZEN BEEF (BONELESS)", "0201300000")]  # the lost leading zero is back
    assert r.items[0]["unit_price"] == "0.0012" and r.items[1]["amount"] == "12"  # no 0.0012000000000000001
    assert r.document == {"invoice_no": "880101773", "invoice_date": "2026-10-15"}
    assert all(c.ok for c in r.checks), r.checks
    assert r.items[0]["_source"].endswith("Invoice row 19")


def test_old_excel_format():
    r = convert([InputFile(XLS.name, XLS.read_bytes())])
    assert _dfe(r) == [("SR-100", "USB CHARGER 5V", "8504405500"), ("SR-200", "POWER CORD 1.8M", "8544.42.9090"),
                       ("SR-300", "SWITCH 2 WAY", "8536500000")]
    assert r.document == {"invoice_no": "SR-2026-118", "invoice_date": "2026-10-09"}
    assert all(c.ok for c in r.checks), r.checks


def test_excel_preview_draws_the_sheet():
    files = [InputFile("CO-AX invoice.xlsx", invoice_xlsx())]
    r = convert(files)
    pv = preview.build(r, files)
    [page] = pv["pages"]
    assert page["source"] == "CO-AX invoice.xlsx / Invoice" and page["image"].startswith("data:image/svg+xml;base64,")
    svg = base64.b64decode(page["image"].split(",", 1)[1]).decode()
    assert "MLG0603P2N2BTZ10" in svg and "电容 CAPACITOR" in svg and ">19<" in svg  # the cells and the Excel row numbers
    assert len(pv["lines"]) == 3 and all(line["page"] == 0 and line["bbox"] for line in pv["lines"])
    assert pv["lines"][2]["bbox"][3] - pv["lines"][2]["bbox"][1] > pv["lines"][0]["bbox"][3] - pv["lines"][0]["bbox"][1]  # with its wrapped row
    assert pv["totals"]["bbox"] is not None and pv["totals_values"]["amount"] == "856"


SIZE_ROWS_HEAD = ["Boxes", "Acme Code", "Description", "HTS CODE", "Fabric", "Fabric & Color Code", "Size", "Pcs", "USD", "EXW/USD"]
SIZE_ROWS = [  # from row 7
    ["Marks & Nos/", None, "No & kind of pkgs.", None, None, "Description of goods", None, "Quantity", "Rate", "Amount"],
    SIZE_ROWS_HEAD,
    ["100 % COTTON LADIES GARMENTS (Powerloom Items)"],
    [12, "10018LW(a)", "Ladies long dress", 6204423020, "Cotton voile", "CORG/22/10", "S", 128, 31.65, 4051.2],  # merged over 3 rows
    [None] * 6 + ["M", 55, 31.65, 1740.75],
    [None] * 6 + ["L", 37, 31.65, 1171.05],
    [],
    SIZE_ROWS_HEAD,  # repeated
    ["100 % COTTON KIDS GARMENTS (Powerloom Items)"],
    [3, "10059KW(a)", "Girls dress", 6204423020, "Gingham", "CORG/054/01", "2-3Y", 18, 27.75, 499.5],  # merged over 2 rows
    [None] * 6 + ["7-8Y", 4, 29.65, 118.6],
    [],
    [2, "8413B(d)", "Hottie bag", 4202224020, "Gingham", "CORG/055/02", "OS", 37, 6.95, 257.15],
    *[[]] * 8,
    ["Amount in words", None, None, None, None, None, "Total Pcs : ", 279, None, 7838.25],
    [],
    ["Total EXW Value USD : Seven thousand eight hundred thirty eight and cents twenty five only."],
]
SIZE_ROWS_MERGED = [(10, 12), (16, 17)]  # rows of the merged product cells (columns B-F)
SIZE_ROWS_XLS = Path(__file__).resolve().parent / "fixtures" / "size_rows.xls"  # the same sheet, old Excel format


def size_rows_xlsx() -> bytes:
    """An export invoice laid out like many Indian exporters' (e.g. Acus Impex): the header written over two
    rows, each product's code, description and HTS code in merged cells next to its size rows, section titles,
    the header repeated, the invoice number under its label, the totals row a few empty rows below with
    "Total Pcs :" in the size column – and a tax invoice sheet (no HTS codes, a tax "Rate") that is not the one."""
    wb = Workbook()
    inv = wb.active
    inv.title = "Invoice "
    tax = wb.create_sheet("Tax Invoice")
    for ws, title in ((inv, "EXPORT INVOICE"), (tax, "TAX INVOICE")):
        ws["A1"], ws["A2"], ws["A3"] = title, "Exporter:", "ACME EXPORTS PVT. LTD."
        ws["F2"], ws["F3"] = "Invoice no & date", "AE-017-26-27 Dt. 28.09.2026"
        ws["F4"], ws["F5"] = "Buyer's code no & date", "Order No. 021A-26-27"
    inv.merge_cells("A1:J1")
    tax_rows = [
        [None, None, None, None, "Description of goods", None, "Quantity", "Rate", "Amount", None, "IGST (INR)"],
        ["Boxes", "Acme Code", "Description", "Fabric", "Fabric & Color Code", "Size", "Pcs", "USD", "EXW / USD", "Taxable Value (INR)", "Rate", "Amount"],
        [12, "10018LW(a)", "Ladies long dress", "Cotton voile", "CORG/22/10", "S", 128, 31.65, 4051.2, 385269.12, 0.18, 69348.44],
    ]
    for ws, rows in ((inv, SIZE_ROWS), (tax, tax_rows)):
        for r, row in enumerate(rows, 7):
            for c, v in enumerate(row, 1):
                ws.cell(row=r, column=c, value=v)
    for first, last in SIZE_ROWS_MERGED:
        for col in "BCDEF":
            inv.merge_cells(f"{col}{first}:{col}{last}")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SIZE_ROWS_LINES = [("10018LW(a)", "Ladies long dress", "6204423020")] * 3 + [("10059KW(a)", "Girls dress", "6204423020")] * 2 + [
    ("8413B(d)", "Hottie bag", "4202224020")]


def test_two_header_rows_and_merged_size_rows():
    files = [InputFile("AE-017-26-27.xlsx", size_rows_xlsx())]
    r = convert(files)
    assert r.customer.id == "auto"
    assert _dfe(r) == SIZE_ROWS_LINES  # every size row is a line, with its product's merged code, description, HTS
    assert [(i["qty"], i["unit_price"], i["amount"]) for i in r.items][3:5] == [("18", "27.75", "499.5"), ("4", "29.65", "118.6")]
    assert [i["_source"] for i in r.items] == [f"AE-017-26-27.xlsx / Invoice  row {n}" for n in (10, 11, 12, 16, 17, 19)]
    assert r.parts[0].labels == {"hs_code": "HTS CODE", "model": "Acme Code", "description": "Description",
                                 "qty": "Quantity Pcs", "unit_price": "Rate USD", "amount": "Amount EXW/USD"}
    assert r.document == {"invoice_no": "AE-017-26-27", "invoice_date": "28.09.2026"}  # under "Invoice no & date"
    assert {c.name: c.ok for c in r.checks} == {"Line items": True, "Subtotal qty": True, "Subtotal amount": True, "Line math": True}
    assert r.parts[0].totals["qty"] == "279" and r.warnings == [] and not r.needs_review
    rows = r.parts[0].layout["rows"]
    assert [(row["kind"], row["excel_row"]) for row in rows] == [("head", 7), ("head", 8)] + [
        ("item", n) for n in (10, 11, 12, 16, 17, 19)] + [("total", 28)]
    pv = preview.build(r, files)
    assert len(pv["lines"]) == 6 and all(line["page"] == 0 and line["bbox"] for line in pv["lines"])
    assert pv["totals"]["bbox"] is not None and pv["totals_values"]["amount"] == "7838.25"


def test_two_header_rows_and_merged_size_rows_old_excel_format():
    r = convert([InputFile(SIZE_ROWS_XLS.name, SIZE_ROWS_XLS.read_bytes())])
    assert _dfe(r) == SIZE_ROWS_LINES
    assert r.document == {"invoice_no": "AE-017-26-27", "invoice_date": "28.09.2026"}
    assert all(c.ok for c in r.checks), r.checks


def test_unknown_supplier_is_read_like_a_known_one():
    """The same ruled table with another seller's name: the general reader gives what the profile gives."""
    known = convert([InputFile("a.pdf", problem_invoice())])
    other = convert([InputFile("a.pdf", problem_invoice(seller="SOMEONE ELSE LTD"))])
    assert (known.customer.id, other.customer.id) == ("istone", "auto")
    assert _dfe(other) == _dfe(known)
    assert [(c.name, c.ok) for c in other.checks] == [(c.name, c.ok) for c in known.checks]
    assert sorted((f.line, f.code) for f in other.flags) == sorted((f.line, f.code) for f in known.flags)


def test_a_known_customer_with_a_new_layout_is_still_read():
    lines = [[f"P-{n}", f"Part {n}", str(n), "2.00", f"{2 * n:.2f}", "9006910001"] for n in range(1, 4)]
    pdf = invoice_pdf(lines, ["Subtotal", "", "6", "", "12.00", ""],
                      header=["Part No.", "Description", "QTY(pcs)", "Unit Price(USD)", "Total (USD)", "HS Code"])
    r = convert([InputFile("istone new layout.pdf", pdf)])
    assert r.customer.id == "istone"  # still recognised by its name; the table was found by the general reader
    assert _dfe(r) == [("P-1", "Part 1", "9006910001"), ("P-2", "Part 2", "9006910001"), ("P-3", "Part 3", "9006910001")]
    assert any("general reader" in w for w in r.warnings)  # the profile did not fit: the document is checked by a person


def test_no_item_table_is_a_clear_error():
    pdf = invoice_pdf([["a", "b", "c", "d", "e", "f"]], ["", "", "", "", "", ""], seller="NOBODY",
                      header=["Name", "Address", "Phone", "Fax", "Mail", "Web"])
    with pytest.raises(ConversionError, match="no line-item table was found"):
        convert([InputFile("letter.pdf", pdf)])


@pytest.mark.parametrize("labels, expected", [
    (["No.", "Part No.", "Description of Goods", "H.S. Code", "Qty", "Unit Price (USD)", "Amount (USD)"],
     {"line_no": 0, "model": 1, "description": 2, "hs_code": 3, "qty": 4, "unit_price": 5, "amount": 6}),
    (["PO#", "Item#", "Description", "HTS", "HTS Description", "Qty", "Unit Price", "Amount"],
     {"model": 1, "description": 2, "hs_code": 3, "qty": 5, "unit_price": 6, "amount": 7}),
    (["ITEM", "Model", "HS CODE", "Ordering Code", "PRODUCT DESCRIPTION（EN）", "QUANTITY\nPCS", "UNIT PRICE\nUSD", "NET AMOUNT\nUSD"],
     {"line_no": 0, "model": 1, "hs_code": 2, "description": 4, "qty": 5, "unit_price": 6, "amount": 7}),
    (["Model", "Description", "QTY(pcs)", "Unit Price(USD)", "Total (USD)", "United States HS Code"],
     {"model": 0, "description": 1, "qty": 2, "unit_price": 3, "amount": 4, "hs_code": 5}),
    (["序号", "型号", "品名", "海关编码", "数量", "单价", "金额"],
     {"line_no": 0, "model": 1, "description": 2, "hs_code": 3, "qty": 4, "unit_price": 5, "amount": 6}),
    (["PO#", "Item#", "Material Description", "Qty / CTN", "Qty", "Total Gross Weight"],
     {"model": 1, "description": 2, "qty": 4, "gross_weight": 5}),
    (["Boxes", "Acus Code", "Description", "HTS CODE", "Fabric", "Fabric & Color Code", "Size", "Pcs", "Rate", "Amount"],
     {"model": 1, "description": 2, "hs_code": 3, "qty": 7, "unit_price": 8, "amount": 9}),
    (["Style", "Description", "HSN Code", "Qty", "IGST Rate", "Rate/Pc", "Amount"],
     {"model": 0, "description": 1, "hs_code": 2, "qty": 3, "unit_price": 5, "amount": 6}),
])
def test_header_wordings(labels, expected):
    m = generic.match_labels(labels, load_generic())
    assert m is not None
    assert {f: i for f, i in m.fields.items() if f in expected} == expected


def test_codes_that_are_not_the_product_code():
    for label in ("Color Code", "Fabric Code", "Country Code", "Bar Code", "HSN Code"):
        m = generic.match_labels([label, "Item Description", "HS Code", "Qty", "Amount"], load_generic())
        assert m is not None and "model" not in m.fields, label


def test_not_a_header():
    assert generic.match_labels(["Name", "Address", "Phone", "Description"], load_generic()) is None


FUYAO_LINES = [
    ["7475833-08", "Automotive Tempered Glass HS:7007.11.0010", "1190*760*780mm", "180", "PCS", "", "DA072239", "USD 6.70", "USD 1206.00"],
    ["8496037-07", "Automotive Laminated Glass (with bracket) HS: 8708.22.0000", "1190*760*780mm", "420", "PCS", "", "DA072239", "USD 17.56", "USD 7375.20"],
    ["4A0C483-08", "Automotive Laminated Glass (with heating line) HS: 8708.22.0000", "1720*1120*1220mm", "180", "PCS", "", "DA072239", "USD 46.12", "USD 8301.60"],
]


def test_headers_that_do_not_say_what_is_under_them():
    """Fuyao Glass: the part numbers under "Description", the description with "HS:…" in it under "Hs Code",
    header words broken over two lines ("Quantit" / "y"), empty ruled rows, a "total" row with freight / FOB /
    insurance in its other cells."""
    r = convert([InputFile("FP-SYFY-G06.pdf", offset_invoice_pdf(FUYAO_LINES))])
    assert _dfe(r) == [("7475833-08", "Automotive Tempered Glass", "7007.11.0010"),
                       ("8496037-07", "Automotive Laminated Glass (with bracket)", "8708.22.0000"),
                       ("4A0C483-08", "Automotive Laminated Glass (with heating line)", "8708.22.0000")]
    assert {c.name: c.ok for c in r.checks} == {"Line items": True, "Subtotal qty": True, "Subtotal amount": True}
    assert r.flags == [] and r.document == {"invoice_no": "SYFY-G06-2026-7-15-1", "invoice_date": "2026-07-15"}
    assert r.parts[0].labels["model"] == "Description" and r.parts[0].labels["qty"] == "Quantit y"
    box = r.parts[0].rows[0].cells  # the preview marks the part number as the product code, the text as the description
    assert box["model"][0] < box["description"][0] and box["hs_code"] == box["description"]


def test_an_hs_code_inside_the_description():
    """No HS code column: "… HS CODE: 9405.40.8440" in the description becomes the HS code. A header word broken
    over two lines without ruling lines ("Quantit" / "y") is read too."""
    ops = ["0.5 w", _text(220, 800, "COMMERCIAL INVOICE", 12), _text(40, 770, "Invoice No.: LT-2026-31")]
    for label, x in (("Item No.", 40), ("Description", 120), ("Quantit", 380), ("Unit Price", 440), ("Amount", 520)):
        ops.append(_text(x, 740, label))
    ops.append(_text(380, 731, "y"))
    lines = [("LP-600", "LED PANEL LIGHT 600X600 HS CODE: 9405.40.8440", "20", "12.50", "250.00"),
             ("LP-300", "LED PANEL LIGHT 300X300 HS CODE: 9405.40.8440", "10", "8.00", "80.00")]
    y = 712
    for item, desc, qty, price, amount in lines:
        ops += [_text(40, y, item), _text(120, y, desc), _right(410, y, qty), _right(490, y, price), _right(560, y, amount)]
        y -= 16
    ops += [_text(120, y - 4, "TOTAL"), _right(410, y - 4, "30"), _right(560, y - 4, "330.00")]
    r = convert([InputFile("lights.pdf", _document(["\n".join(ops).encode("latin-1")]))])
    assert _dfe(r) == [("LP-600", "LED PANEL LIGHT 600X600", "9405.40.8440"), ("LP-300", "LED PANEL LIGHT 300X300", "9405.40.8440")]
    assert all(c.ok for c in r.checks) and {c.name for c in r.checks} >= {"Subtotal qty", "Line math"}


AMPHENOL_COLUMNS = [("SN", 40), ("Part Number", 80), ("Coded Part", 180), ("Invoice No.", 280),
                    ("Quantity", 360), ("Unit Price", 420), ("Total Value", 490),
                    ("Description Of Goods", 580), ("HS Code", 700)]


def amphenol_pdf() -> bytes:
    """An Amphenol-style packing list & invoice (landscape, no ruling lines): one item whose cells are
    written over three lines – the description stands above *and* below the line with the numbers – and
    notes under the item ("Lot No …", "Material …") that are not part of it."""
    ops = ["0.5 w", _text(300, 560, "PACKING LIST & COMMERCIAL INVOICE", 12)]
    ops += [_text(x, 520, label) for label, x in AMPHENOL_COLUMNS]
    ops += [_text(180, 500, "Cable assembly"), _text(580, 500, "Insulated wire")]  # above the numbers
    for x, value in ((40, "1"), (80, "CWR10-020-09ST"), (280, "10939717"), (360, "19200"),
                     (420, "1.317427"), (490, "25294.5984"), (700, "8544492100")):
        ops.append(_text(x, 495, value))
    ops += [_text(180, 490, "without connector"), _text(580, 490, "300V RADSOK brand")]  # below the numbers
    ops += [_text(40, 476, "Lot No: 5059896501"), _text(180, 464, "Material Copper Rubber")]
    return _document(["\n".join(ops).encode("latin-1")], width=842, height=595)


def test_a_cell_written_above_and_below_the_line_with_the_numbers():
    """Amphenol: the description cell has two lines while the rest of its row has one, so half of it
    stands above the numbers and half below – lines that overlap vertically are one row. The notes
    under the item are not wrapped text of it, and the invoice number is the one in the table (the
    text says "Invoice No." only as a column header, with "Quantity" behind it)."""
    r = convert([InputFile("26090643.pdf", amphenol_pdf())])
    assert _dfe(r) == [("CWR10-020-09ST", "Insulated wire 300V RADSOK brand", "8544492100")]
    assert r.document == {"invoice_no": "10939717"}
    assert "Copper" not in str(r.items) and "Lot No" not in str(r.items)
    assert r.items[0]["qty"] == "19200" and r.items[0]["unit_price"] == "1.317427"
    assert [c.ok for c in r.checks if c.name == "Line math"] == [True]


def _chinese_part(description: str, others: dict[str, str]) -> PdfExtract:
    box = (0.0, 0.0, 1.0, 1.0)
    return PdfExtract(name="x.pdf", labels={"description": "Description Of Goods"},
                      items=[{"model": "CWR10-020-09ST", "description": description}],
                      rows=[RowBox(1, box, {"description": box})],
                      others=[{k: (v, (2.0, 0.0, 3.0, 1.0)) for k, v in others.items()}])


def test_a_description_column_of_only_chinese_takes_the_english_of_another_column():
    """Amphenol: "Description Of Goods" is the Chinese customs name 中国海关品名 and the English
    description stands under "Coded Part" – it is used instead, and the line is marked for review."""
    part = _chinese_part("电缆/无接头 /300V/RADSOK牌",
                         {"Coded Part": "UL3132 10AWG wire, BK", "Currency": "USD", "Remark": "600 PCS*32"})
    generic.english_description(part)
    assert part.items[0]["description"] == "UL3132 10AWG wire, BK"
    assert part.rows[0].cells["description"] == (2.0, 0.0, 3.0, 1.0)  # the preview marks the column it came from
    assert len(part.notes) == 1 and "Coded Part" in part.notes[0]


@pytest.mark.parametrize("description, others", [
    ("Pivot Stand 转向支架", {"Coded Part": "UL3132 10AWG wire, BK"}),  # English of its own: left alone
    ("Automotive Tempered Glass", {"Coded Part": "UL3132 10AWG wire, BK"}),
    ("电缆/无接头 /300V/RADSOK牌", {"Coded Part": "DXCWR10-020-01", "Units": "M"}),  # no English to take
    ("电缆/无接头 /300V/RADSOK牌", {"Coded Part": "电缆 无接头"}),
])
def test_the_description_of_other_invoices_is_left_alone(description, others):
    part = _chinese_part(description, others)
    generic.english_description(part)
    assert part.items[0]["description"] == description and part.notes == []


def test_a_proper_hs_code_column_is_left_alone():
    r = convert([InputFile("u.pdf", problem_invoice(seller="SOMEONE ELSE LTD"))])
    assert r.rows[0][5] == "9006910001" and r.rows[0][4] == "Test Product 1"


@pytest.mark.parametrize("title, what", [
    ("ARRIVAL NOTICE   Type of B/L: Sea WayBill   B/L Number: COSU6506620260", "an Arrival Notice"),
    ("DECLARATION  Importation of Motor Vehicles and Motor Vehicle Equipment Subject to Federal Motor Vehicle Safety", "an HS-7 declaration"),
    ("BILL OF LADING FOR OCEAN TRANSPORT", "a bill of lading"),
    ("CERTIFICATE OF ORIGIN", "a certificate of origin"),
])
def test_documents_that_are_not_invoices_say_what_they_are(title, what):
    ops = [_text(40, 800, title, 10), _text(40, 760, "CONTAINER NO. TGBU4887141   40HQ   51 Packages   16115.62 KG"),
           _text(40, 745, "AUTOMOTIVE LAMINATED GLASS (WITH BRACKET)")]
    with pytest.raises(ConversionError, match=f"this is {re.escape(what)}.*not a Commercial Invoice"):
        convert([InputFile("notice.pdf", _document(["\n".join(ops).encode("latin-1")]))])
