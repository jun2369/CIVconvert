import base64
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from civconvert import preview
from civconvert.api import app
from civconvert.pipeline import InputFile, convert
from conftest import SAMPLE, needs_sample
from pdfgen import problem_invoice

client = TestClient(app)


def _inside(inner, outer, slack=0.5):
    return inner[0] >= outer[0] - slack and inner[1] >= outer[1] - slack and inner[2] <= outer[2] + slack and inner[3] <= outer[3] + slack


def test_multi_page_invoice_lines_are_placed_on_their_pages():
    files = [InputFile("problem.pdf", problem_invoice())]
    r = convert(files, customer="istone")
    data = preview.build(r, files)

    assert [(p["source"], p["page"]) for p in data["pages"]] == [("problem.pdf", 1), ("problem.pdf", 2)]
    assert all(p["image"].startswith("data:image/") for p in data["pages"])
    assert len(data["lines"]) == len(r.items) == 26
    assert [line["page"] for line in data["lines"]] == [0] * 14 + [1] * 12
    for line in data["lines"]:  # every line and its cells are inside the table cut out of its page
        table = data["pages"][line["page"]]["table"]
        assert _inside(line["bbox"], table)
        assert set(line["cells"]) == {"model", "description", "qty", "unit_price", "amount", "hs_code"}
        assert all(_inside(c, line["bbox"]) for c in line["cells"].values())
    tops = [line["bbox"][1] for line in data["lines"][:14]]
    assert tops == sorted(tops)
    # the header is repeated on page 2; the subtotal row is on page 2
    assert [h["page"] for h in data["headers"]] == [0, 1]
    assert data["totals"]["page"] == 1 and _inside(data["totals"]["bbox"], data["pages"][1]["table"])
    assert data["fields"][0] == {"name": "model", "label": "Model"}


def test_problems_are_flagged_on_their_cells():
    r = convert([InputFile("problem.pdf", problem_invoice())], customer="istone")
    flags = {(f.line, f.level, tuple(f.fields), tuple(f.columns)) for f in r.flags}
    assert (5, "warning", ("hs_code",), ("F",)) in flags  # unexpected HS code format
    assert (9, "error", ("qty", "unit_price", "amount"), ()) in flags  # qty x price != amount
    assert (12, "error", ("description",), ("E",)) in flags  # empty description
    assert (3, "warning", ("model", "hs_code"), ("D", "F")) in flags  # same model, other HS code on line 20
    assert (20, "warning", ("model", "hs_code"), ("D", "F")) in flags
    assert {f.line for f in r.flags} == {3, 5, 9, 12, 20}
    # the subtotal still adds up, line math does not
    checks = {c.name: c.ok for c in r.checks}
    assert checks["Subtotal qty"] and checks["Subtotal amount"] and not checks["Line math"]
    assert "Line 20: model 'CINSDUP(X1)' has hs_code '8518108030' here but '9006910001' on line 3" in r.warnings
    assert r.to_dict()["flags"][0]["line"] == 3


def test_grid_is_read_back_from_the_exported_files():
    r = convert([InputFile("problem.pdf", problem_invoice())], customer="istone")
    grid = preview.grid(r)
    assert [c["letter"] for c in grid["columns"]] == ["A", "D", "E", "F"]
    assert [c["header"] for c in grid["columns"]] == ["LNO", "Product Code", "Goods Description", "Tariff"]
    assert grid["total_columns"] == 210 and grid["header_row"] == 1
    lno, model, desc, tariff = grid["columns"]
    assert lno["cells"][0] == {"v": "1", "num": True} and model["cells"][0] == {"v": "CINSM001"}
    assert desc["cells"][11] == {"v": ""}  # line 12 has no description
    assert lno["head_style"]["bold"] and model["width"] > lno["width"]
    # the same values as the Excel file the user downloads
    ws = load_workbook(io.BytesIO(r.xlsx_bytes())).worksheets[0]
    assert [ws.cell(row=n + 2, column=6).value for n in range(26)] == [c["v"] for c in tariff["cells"]]


@needs_sample
def test_sample_preview_cuts_out_the_table():
    files = [InputFile(SAMPLE.name, SAMPLE.read_bytes())]
    r = convert(files)
    data = preview.build(r, files)
    [page] = data["pages"]
    x0, top, x1, bottom = page["table"]
    assert top < data["headers"][0]["bbox"][1] and bottom > data["totals"]["bbox"][3]
    assert 0 <= x0 < 51 and 541 < x1 <= page["width"]
    assert len(data["lines"]) == 20 and r.flags == []
    assert data["totals_values"]["amount"] == "US$48,561.00"


@needs_sample
def test_preview_endpoint(api_key):
    res = client.post("/api/preview", files={"files": (SAMPLE.name, SAMPLE.read_bytes(), "application/pdf")}, headers=api_key)
    assert res.status_code == 200
    body = res.json()
    assert body["result"]["filename"] == "CargoWise Input - 8100800394 Commercial Invoice.csv"
    assert base64.b64decode(body["result"]["content_base64"]).startswith(b"LNO,")
    assert body["result"]["xlsx_base64"] and len(body["preview"]["pages"]) == 1


def test_preview_endpoint_reports_conversion_errors(api_key):
    res = client.post("/api/preview", files={"files": ("x.pdf", b"hello", "application/pdf")}, headers=api_key)
    assert res.status_code == 422 and "not a PDF" in res.json()["error"]


def test_page_and_scripts_are_served():
    assert "CIVPreview" in client.get("/static/preview.js").text
    assert client.get("/static/preview.css").status_code == 200
    page = client.get("/login")
    assert "/static/app.js" in page.text and page.headers["cache-control"] == "no-cache"
