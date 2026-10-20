"""Write the filled template as CSV (the CargoWise import file) or XLSX (for review)."""

from __future__ import annotations

import csv
import io
import re

from openpyxl import load_workbook

from .template import Template

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def csv_headers(template: Template) -> list[str]:
    if template.profile.header_control_chars == "strip":
        return [_CONTROL.sub("", h).rstrip() for h in template.headers]
    return list(template.headers)


def to_csv(template: Template, rows: list[list[str]]) -> bytes:
    p = template.profile
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, delimiter=p.csv_delimiter, lineterminator=p.csv_line_terminator)
    writer.writerow(csv_headers(template))
    writer.writerows(rows)
    return buf.getvalue().encode(p.csv_encoding, errors="replace")


def _number(value: str) -> int | float | str:
    for kind in (int, float):
        try:
            return kind(value)
        except ValueError:
            pass
    return value


def to_xlsx(template: Template, rows: list[list[str]], numeric: frozenset[int] = frozenset()) -> bytes:
    """The template workbook (styles kept) with its sample rows replaced by `rows`."""
    p = template.profile
    wb = load_workbook(p.workbook)
    ws = wb[p.sheet] if p.sheet else wb.worksheets[0]
    first = p.header_row + 1
    if ws.max_row >= first:
        ws.delete_rows(first, ws.max_row - first + 1)
    for r, row in enumerate(rows, start=first):
        for c, value in enumerate(row):
            if value != "":
                ws.cell(row=r, column=c + 1, value=_number(value) if c in numeric else value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
