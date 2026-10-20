"""The CargoWise template workbook: its header row defines the output columns."""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from .profiles import ConfigError, TemplateProfile

_OOXML_ESCAPE = re.compile(r"_x([0-9A-Fa-f]{4})_")
_LETTERS = re.compile(r"[A-Za-z]{1,3}")


def unescape_ooxml(text: str) -> str:
    """'_x001F_' -> '\\x1f', the way Excel itself reads these escapes."""
    return _OOXML_ESCAPE.sub(lambda m: chr(int(m.group(1), 16)), text)


def normalize_header(text: str) -> str:
    return " ".join(re.sub(r"[\x00-\x1f\x7f]", " ", text).split()).lower()


@dataclass
class Template:
    profile: TemplateProfile
    headers: list[str]

    @property
    def id(self) -> str:
        return self.profile.id

    def resolve(self, ref: str) -> int:
        """Column letter ("D") or header text ("Product Code") -> 0-based column index."""
        ref = ref.strip()
        if _LETTERS.fullmatch(ref):
            try:
                idx = column_index_from_string(ref.upper()) - 1
            except ValueError:
                idx = len(self.headers)
            if idx < len(self.headers):
                return idx
        wanted = normalize_header(ref)
        hits = [i for i, h in enumerate(self.headers) if normalize_header(h) == wanted]
        if len(hits) == 1:
            return hits[0]
        if hits:
            raise ConfigError(f"template '{self.id}' has {len(hits)} columns named '{ref}' – use the column letter instead")
        last = get_column_letter(len(self.headers))
        raise ConfigError(f"template '{self.id}' has no column '{ref}' (its columns are A-{last})")


def load_template(profile: TemplateProfile) -> Template:
    if not profile.workbook.is_file():
        raise ConfigError(f"template workbook not found: {profile.workbook}")
    wb = load_workbook(profile.workbook, read_only=True, data_only=True)
    try:
        ws = wb[profile.sheet] if profile.sheet else wb.worksheets[0]
        rows = ws.iter_rows(min_row=profile.header_row, max_row=profile.header_row, values_only=True)
        header = next(rows, ())
    except KeyError as e:
        raise ConfigError(f"template '{profile.id}': sheet '{profile.sheet}' not found") from e
    finally:
        wb.close()
    headers = ["" if v is None else unescape_ooxml(str(v)) for v in header]
    while headers and not headers[-1]:
        headers.pop()
    if not headers:
        raise ConfigError(f"template '{profile.id}': row {profile.header_row} of {profile.workbook.name} is empty")
    return Template(profile, headers)


# ───── how a workbook shows cells (the preview and the online sheet look like the exported Excel) ─────


def shown_value(value: Any) -> str:
    """A cell value the way Excel shows it with the General number format."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return format(value, ".11g")
    return str(value)


def _rgb(color) -> str | None:
    if color is None or color.type != "rgb" or not isinstance(color.rgb, str):
        return None  # theme / indexed colors: leave it to the default
    return "#" + color.rgb[-6:]


def cell_style(cell) -> dict[str, Any]:
    font, fill = cell.font, cell.fill
    style: dict[str, Any] = {"font": font.name or "Arial", "size": float(font.sz or 10), "bold": bool(font.b)}
    if color := _rgb(font.color):
        style["color"] = color
    if fill is not None and fill.fill_type == "solid" and (bg := _rgb(fill.fgColor)):
        style["fill"] = bg
    return style


def column_width_px(ws, col: int) -> int:
    """Column width in screen pixels (Excel: characters of the default font, ~7 px each + 5 px padding)."""
    width = ws.sheet_format.defaultColWidth or 8.43
    for dim in ws.column_dimensions.values():
        if dim.min and dim.max and dim.min <= col <= dim.max and dim.width:
            width = dim.width
            break
    return int(width * 7 + 5)
