"""What kind of file an invoice is – PDF, Excel or a picture – and Excel workbooks as plain grids of cell values.

The file type is told by the content, not only by the name: a PDF starts with %PDF, an .xlsx is a zip
holding xl/workbook.xml, an .xls is an OLE2 file, a picture (PNG, JPEG, TIFF, BMP, GIF, WEBP, HEIC, AVIF)
has its format's signature. Cell values become text the way Excel shows them in the General format
(no 0.0012000000000000001), dates as YYYY-MM-DD. A picture is read as a PDF with one page per picture
page (picture_pdf); OCR (ocr.py) reads its text.
"""

from __future__ import annotations

import datetime as dt
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from typing import Any

from .formula import format_number

PDF, XLSX, XLS, IMAGE = "pdf", "xlsx", "xls", "image"
CONTENT_TYPES = {
    PDF: "application/pdf",
    XLSX: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    XLS: "application/vnd.ms-excel",
}
PICTURE_TYPES = {
    "png": "image/png", "jpeg": "image/jpeg", "tiff": "image/tiff", "bmp": "image/bmp", "gif": "image/gif",
    "webp": "image/webp", "heic": "image/heic", "avif": "image/avif",
}
EXTENSIONS = (".pdf", ".xlsx", ".xlsm", ".xls", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp", ".heic", ".heif", ".avif")
MAX_ROWS, MAX_COLS = 3000, 80  # of a sheet that are read
MAX_PICTURE_PAGES = 20  # of a multi-page TIFF
MAX_PICTURE_SIDE = 4000  # px: larger pictures are made smaller (OCR reads at most ~3200 px anyway)
_OLE2 = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_HEIF_BRANDS = (b"heic", b"heix", b"hevc", b"hevx", b"heim", b"heis", b"mif1", b"msf1")
_ZERO_WIDTH = re.compile(r"[​-‏⁠﻿­]")
_SPACES = re.compile(r"\s+")


class SourceError(Exception):
    """The file cannot be read."""


def kind(name: str, data: bytes) -> str | None:
    """PDF, XLSX, XLS, IMAGE – or None for anything else."""
    if b"%PDF" in data[:1024]:
        return PDF
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "xl/workbook.xml" in z.namelist():
                    return XLSX
        except zipfile.BadZipFile:
            return None
        return None
    if data[:8] == _OLE2 and name.lower().endswith(".xls"):
        return XLS
    if picture_format(name, data):
        return IMAGE
    return None


def picture_format(name: str, data: bytes) -> str | None:
    """png, jpeg, tiff, bmp, gif, webp, heic (iPhone photos), avif – by the file's signature – or None."""
    head = data[:16]
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if head.startswith(b"\xff\xd8\xff"):
        return "jpeg"
    if head[:4] in (b"II*\x00", b"MM\x00*", b"II+\x00", b"MM\x00+"):
        return "tiff"
    if head.startswith((b"GIF87a", b"GIF89a")):
        return "gif"
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return "webp"
    if head[4:8] == b"ftyp":
        if head[8:12] in (b"avif", b"avis"):
            return "avif"
        if head[8:12] in _HEIF_BRANDS:
            return "heic"
    if head.startswith(b"BM") and name.lower().endswith(".bmp"):  # two bytes say little: the name must agree
        return "bmp"
    return None


def content_type(name: str, data: bytes) -> str:
    k = kind(name, data)
    if k == IMAGE:
        return PICTURE_TYPES[picture_format(name, data) or "png"]
    return CONTENT_TYPES[k or PDF]


def strip_extension(name: str) -> str:
    return re.sub(r"\.(pdf|xlsx|xlsm|xls|png|jpe?g|tiff?|bmp|gif|webp|heic|heif|avif)$", "", name, flags=re.IGNORECASE)


# ───────────────────────────── pictures ─────────────────────────────


def _heif_opener() -> None:
    """HEIC / HEIF (iPhone photos) need the pillow-heif plugin; without it they cannot be opened."""
    try:
        from pillow_heif import register_heif_opener
    except ImportError:
        return
    register_heif_opener()


def _upright(image):
    """RGB (or grey) on white, turned as the camera held it, at most MAX_PICTURE_SIDE px."""
    from PIL import Image, ImageOps

    image = ImageOps.exif_transpose(image)
    if image.mode in ("RGBA", "LA", "PA") or (image.mode == "P" and "transparency" in image.info):
        rgba = image.convert("RGBA")
        white = Image.new("RGB", rgba.size, "white")
        white.paste(rgba, mask=rgba.getchannel("A"))
        image = white
    elif image.mode in ("I;16", "I;16B", "I;16L", "I", "F"):  # 16-bit scans
        image = image.convert("F")
        top = image.getextrema()[1] or 1
        scale = 255.0 / top if top > 255 else 1.0
        image = image.point(lambda v: v * scale).convert("L")
    elif image.mode not in ("RGB", "L"):
        image = image.convert("RGB")
    if max(image.size) > MAX_PICTURE_SIDE:
        image = image.copy()
        image.thumbnail((MAX_PICTURE_SIDE, MAX_PICTURE_SIDE), Image.LANCZOS)
    return image


def picture_pages(data: bytes, name: str) -> list[Any]:
    """The pages of a picture as PIL images (a TIFF may hold several pages; of an animation the first picture)."""
    from PIL import Image, ImageSequence

    _heif_opener()
    fmt = picture_format(name, data)
    try:
        with Image.open(io.BytesIO(data)) as im:
            pages = []
            for frame in ImageSequence.Iterator(im):
                pages.append(_upright(frame.copy()))
                if fmt != "tiff" or len(pages) >= MAX_PICTURE_PAGES:
                    break
    except Exception as e:  # Pillow raises many different exception types
        raise SourceError(f"{name}: cannot read the picture ({e.__class__.__name__}: {e})") from e
    if not pages or min(pages[0].size) < 8:
        raise SourceError(f"{name}: the picture is empty")
    return pages


def page_points(size: tuple[int, int]) -> float:
    """Points per pixel of a picture page: as wide as an A4 page (portrait), or as an A4 page is high (landscape)."""
    w, h = size
    return (842.0 if w > h else 595.0) / w


def picture_pdf(data: bytes, name: str) -> tuple[bytes, list[Any]]:
    """A picture as a PDF – one page per picture page, each the size page_points() gives – and the pages
    themselves (OCR reads their pixels). The PDF is only drawn on screen (the preview): JPEG quality 92."""
    pages = picture_pages(data, name)
    buf = io.BytesIO()
    first = pages[0]
    first.save(buf, "PDF", resolution=72.0 / page_points(first.size), quality=92, save_all=True, append_images=pages[1:])
    return buf.getvalue(), pages


@dataclass
class Grid:
    """One worksheet: rows of cell values (str, int, float, bool, date/datetime or None)."""

    name: str
    cells: list[list[Any]]
    # merged cells: (first row, first column, last row, last column), 0-based; the value is in the first cell
    merged: list[tuple[int, int, int, int]] = field(default_factory=list)

    def text(self) -> str:
        return "\n".join(" ".join(t for t in (cell_text(v) for v in row) if t) for row in self.cells)


def clean_text(value: str) -> str:
    return _SPACES.sub(" ", _ZERO_WIDTH.sub("", value)).strip()


def cell_text(value: Any) -> str:
    """A cell value as Excel's General format shows it."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return format_number(value)
    return clean_text(unicodedata.normalize("NFC", str(value)))


def read_workbook(data: bytes, name: str, file_kind: str) -> list[Grid]:
    """The visible worksheets of an .xlsx / .xlsm / .xls file."""
    try:
        if file_kind == XLSX:
            return _read_xlsx(data)
        if file_kind == XLS:
            return _read_xls(data)
    except SourceError:
        raise
    except Exception as e:  # the readers raise many different exception types
        raise SourceError(f"{name}: cannot read the Excel file ({e.__class__.__name__}: {e})") from e
    raise SourceError(f"{name}: not an Excel file")


def _read_xlsx(data: bytes) -> list[Grid]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)  # formulas: the values Excel saved with them
    grids = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        rows = []
        for row in ws.iter_rows(max_row=min(ws.max_row, MAX_ROWS), max_col=min(ws.max_column, MAX_COLS), values_only=True):
            rows.append(list(row))
        merged = [(m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1) for m in ws.merged_cells.ranges]
        grids.append(Grid(ws.title, rows, [m for m in merged if m[0] < MAX_ROWS and m[1] < MAX_COLS]))
    return grids


def _read_xls(data: bytes) -> list[Grid]:
    import xlrd

    try:  # with the cell formats: needed for the merged cells
        book = xlrd.open_workbook(file_contents=data, formatting_info=True)
    except Exception:
        book = xlrd.open_workbook(file_contents=data)
    grids = []
    for sheet in book.sheets():
        if getattr(sheet, "visibility", 0):
            continue
        rows = []
        for r in range(min(sheet.nrows, MAX_ROWS)):
            row = []
            for c in range(min(sheet.ncols, MAX_COLS)):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_NUMBER:
                    row.append(int(cell.value) if float(cell.value).is_integer() else cell.value)
                elif cell.ctype == xlrd.XL_CELL_TEXT:
                    row.append(cell.value)
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(cell.value))
                else:
                    row.append(None)
            rows.append(row)
        merged = [(r0, c0, r1 - 1, c1 - 1) for r0, r1, c0, c1 in getattr(sheet, "merged_cells", [])]
        grids.append(Grid(sheet.name, rows, [m for m in merged if m[0] < MAX_ROWS and m[1] < MAX_COLS]))
    return grids
