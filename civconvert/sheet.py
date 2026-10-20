"""The editable CargoWise sheet of a document.

rows = [{"id": "r1a2b3", "src": 7, "cells": {"3": {"v": "CINSBAVE"}, "9": {"f": "=ROUND(G2*H2,2)", "v": "12.5"}}}]

- the columns are always the template's columns (headers, order and count never change);
  cells are keyed by the 0-based column index and only non-empty cells are stored
- "v" is the text exactly as typed (HS codes keep their leading zeros); for formula cells
  "f" is the formula and "v" its computed value (recomputed here on every save)
- "src" is the PDF line the row came from (None for a row added by hand), which keeps the
  link to the original invoice after rows are deleted or inserted
"""

from __future__ import annotations

import functools
import re
import uuid
from dataclasses import dataclass, field
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .export import to_csv, to_xlsx
from .formula import ERROR_CODES, recalculate
from .pipeline import OutputColumn, Result, resolve_columns
from .profiles import ColumnSpec, CustomerProfile, load_template_profile
from .template import Template, cell_style, column_width_px, load_template

MAX_ROWS = 5000
MAX_TEXT = 5000
_CONTROL = re.compile(r"[\x00-\x1f\x7f]")


class SheetError(ValueError):
    """The rows sent by the browser are not a valid sheet."""


def new_row_id() -> str:
    return uuid.uuid4().hex[:12]


# ───────────────────────────── the template ─────────────────────────────


@functools.lru_cache(maxsize=16)
def _template_info(template_id: str, workbook_mtime: float) -> dict[str, Any]:
    profile = load_template_profile(template_id)
    template = load_template(profile)
    wb = load_workbook(profile.workbook)
    ws = wb[profile.sheet] if profile.sheet else wb.worksheets[0]
    columns = [
        {
            "index": i,
            "letter": get_column_letter(i + 1),
            "header": " ".join(_CONTROL.sub(" ", h).split()),  # shown without the invisible control characters
            "width": column_width_px(ws, i + 1),
        }
        for i, h in enumerate(template.headers)
    ]
    head = cell_style(ws.cell(row=profile.header_row, column=1))
    data = cell_style(ws.cell(row=ws.max_row + 100, column=1))  # a new cell: the workbook's default font (as in the export)
    wb.close()
    return {"id": template.id, "header_row": profile.header_row, "columns": columns, "head_style": head, "cell_style": data}


def template_info(template_id: str) -> dict[str, Any]:
    """Columns (letter, header, width), header row number and fonts of a template – for the online sheet."""
    workbook = load_template_profile(template_id).workbook
    return _template_info(template_id, workbook.stat().st_mtime)


def headers(template: Template) -> list[str]:
    return list(template.headers)


def mapped_columns(profile: CustomerProfile, template: Template) -> list[OutputColumn]:
    return resolve_columns(profile, template)


def column_roles(columns: list[OutputColumn]) -> dict[str, dict[str, Any]]:
    """letter -> what the customer profile fills into it (field, number, line number)."""
    return {
        c.letter: {"field": c.spec.field, "number": c.spec.number, "line_no": (c.spec.value or "").strip() == "{line_no}"}
        for c in columns
    }


# ───────────────────────────── rows ─────────────────────────────


def rows_from_result(result: Result) -> list[dict[str, Any]]:
    """The converted template rows as sheet rows."""
    return [
        {"id": new_row_id(), "src": n, "cells": {str(c): {"v": v} for c, v in enumerate(row) if v != ""}}
        for n, row in enumerate(result.rows, 1)
    ]


def clean_rows(rows: Any, width: int) -> list[dict[str, Any]]:
    """Validate rows sent by the browser: only cells inside the template, text or formulas, nothing else."""
    if not isinstance(rows, list):
        raise SheetError("rows must be a list")
    if len(rows) > MAX_ROWS:
        raise SheetError(f"too many rows (at most {MAX_ROWS})")
    out, seen = [], set()
    for i, row in enumerate(rows, 1):
        if not isinstance(row, dict):
            raise SheetError(f"row {i} is not an object")
        rid = row.get("id")
        if not isinstance(rid, str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,40}", rid) or rid in seen:
            rid = new_row_id()
        seen.add(rid)
        src = row.get("src")
        src = src if isinstance(src, int) and not isinstance(src, bool) and src > 0 else None
        cells_in = row.get("cells") or {}
        if not isinstance(cells_in, dict):
            raise SheetError(f"row {i}: cells must be an object")
        cells: dict[str, dict[str, str]] = {}
        for key, cell in cells_in.items():
            try:
                col = int(key)
            except (TypeError, ValueError):
                raise SheetError(f"row {i}: '{key}' is not a column index") from None
            if not 0 <= col < width:
                raise SheetError(f"row {i}: column {col + 1} is outside the template ({width} columns)")
            if not isinstance(cell, dict):
                raise SheetError(f"row {i}: cell {get_column_letter(col + 1)} is not an object")
            formula, value = cell.get("f"), cell.get("v")
            if isinstance(formula, str) and formula.startswith("=") and len(formula) > 1:
                cells[str(col)] = {"f": formula[:MAX_TEXT], "v": ""}
            elif value is not None and value != "":
                text = value if isinstance(value, str) else str(value)
                cells[str(col)] = {"v": text[:MAX_TEXT]}
        out.append({"id": rid, "src": src, "cells": cells})
    return out


def compute(rows: list[dict[str, Any]], template: Template) -> list[tuple[int, int, str]]:
    """Recompute every formula; returns the cells whose result is an error."""
    return recalculate(rows, template.profile.header_row, headers(template))


def display_rows(rows: list[dict[str, Any]], width: int) -> list[list[str]]:
    """Full-width rows of text, as written to the CSV (formulas -> their values)."""
    out = []
    for row in rows:
        values = [""] * width
        for key, cell in (row.get("cells") or {}).items():
            col = int(key)
            if 0 <= col < width:
                values[col] = str(cell.get("v", ""))
        out.append(values)
    return out


@dataclass
class SheetExport:
    """A document's sheet as the export files see it (also used by the preview's right-hand side)."""

    template: Template
    columns: list[OutputColumn]  # shown in the preview: the mapped columns and any other column with data
    rows: list[list[str]]
    numeric: frozenset[int] = field(default_factory=frozenset)

    def csv_bytes(self) -> bytes:
        return to_csv(self.template, self.rows)

    def xlsx_bytes(self) -> bytes:
        return to_xlsx(self.template, self.rows, self.numeric)


def export_view(template: Template, mapped: list[OutputColumn], rows: list[dict[str, Any]]) -> SheetExport:
    width = len(template.headers)
    values = display_rows(rows, width)
    columns = list(mapped)
    used = {c.index for c in mapped}
    for col in sorted({int(k) for row in rows for k in (row.get("cells") or {})} - used):
        letter = get_column_letter(col + 1)
        columns.append(OutputColumn(col, letter, template.headers[col], ColumnSpec(letter, None, None, [], None, False)))
    numeric = frozenset(c.index for c in mapped if c.spec.number)
    return SheetExport(template, sorted(columns, key=lambda c: c.index), values, numeric)


# ───────────────────────────── checks ─────────────────────────────


def validate(
    rows: list[dict[str, Any]],
    template: Template,
    profile: CustomerProfile | None,
    mapped: list[OutputColumn],
    extraction: dict[str, Any] | None = None,
) -> list[dict[str, Any]]:
    """Problems in the current sheet, located by row (1-based) and template column letters."""
    flags: list[dict[str, Any]] = []
    header_row = template.profile.header_row

    def flag(line: int, level: str, code: str, message: str, columns: list[str] = (), fields: list[str] = (), **params: Any) -> None:
        # message: English; the web page translates `code` with `params`
        flags.append({"line": line, "level": level, "code": code, "params": params, "message": message,
                      "columns": list(columns), "fields": list(fields)})

    def value(row: dict[str, Any], col: int) -> str:
        return str(((row.get("cells") or {}).get(str(col)) or {}).get("v", ""))

    for n, row in enumerate(rows, 1):
        for key, cell in (row.get("cells") or {}).items():
            if cell.get("f") and cell.get("v") in ERROR_CODES:
                letter = get_column_letter(int(key) + 1)
                ref = f"{letter}{header_row + n}"
                flag(n, "error", "formula_error", f"{ref}: the formula gives {cell['v']} ({cell['f']})", [letter],
                     cell=ref, value=cell["v"], formula=cell["f"])
        if profile is None:
            continue
        for col in mapped:
            f = col.spec.field
            if not f:
                continue
            v = value(row, col.index)
            header = col.header.strip()
            if f in profile.table.required and not v.strip():
                flag(n, "error", "required_empty", f"{col.letter} {header} is empty", [col.letter], [f], letter=col.letter, header=header)
                continue
            rx = profile.checks["patterns"].get(f)
            if rx is not None and v and not rx.search(v):
                flag(n, "warning", "bad_format", f"{col.letter} {header} '{v}' has an unexpected format", [col.letter], [f],
                     letter=col.letter, header=header, value=v)
            if profile.checks["ascii_only"] and v and not v.isascii():
                flag(n, "warning", "non_ascii_cell", f"{col.letter} {header} contains non-English characters: '{v}'", [col.letter], [f],
                     letter=col.letter, header=header, value=v)

    if profile is not None:  # the same model must have the same HS code (checks.consistent)
        by_field = {}
        for col in mapped:
            if col.spec.field:
                by_field.setdefault(col.spec.field, col)
        for key, dependents in profile.checks["consistent"].items():
            key_col = by_field.get(key)
            if key_col is None:
                continue
            groups: dict[str, list[int]] = {}
            for n, row in enumerate(rows, 1):
                k = value(row, key_col.index).strip()
                if k:
                    groups.setdefault(k, []).append(n)
            for dep in dependents:
                dep_col = by_field.get(dep)
                if dep_col is None:
                    continue
                for k, group in groups.items():
                    for n in group:
                        mine = value(rows[n - 1], dep_col.index)
                        others = [{"row": m, "value": value(rows[m - 1], dep_col.index)} for m in group if value(rows[m - 1], dep_col.index) != mine]
                        if others:
                            key_h, dep_h = key_col.header.strip(), dep_col.header.strip()
                            text = ", ".join(f"'{o['value']}' on row {o['row']}" for o in others)
                            flag(n, "warning", "inconsistent_cells", f"{key_h} '{k}' has {dep_h} '{mine}' here, but {text}",
                                 [key_col.letter, dep_col.letter], [key, dep], key_header=key_h, key=k, dep_header=dep_h, mine=mine, others=others)

    # problems that only the PDF shows (qty x price, wrapped rows): stay with the row that came from that line
    if extraction:
        by_src: dict[int, list[int]] = {}
        for n, row in enumerate(rows, 1):
            if row.get("src"):
                by_src.setdefault(row["src"], []).append(n)
        for f in extraction.get("flags") or []:
            if f.get("columns"):
                continue  # about a template cell: checked above on the current value
            inner = {"code": f.get("code", ""), "params": f.get("params") or {}, "message": f.get("message", "")}
            for n in by_src.get(f.get("line"), []):
                flag(n, f.get("level", "warning"), "pdf_line", f"PDF line {f['line']}: {inner['message']}", [], f.get("fields") or [],
                     pdf_line=f["line"], inner=inner)
    flags.sort(key=lambda x: x["line"])
    return flags


def needs_review(flags: list[dict[str, Any]], extraction: dict[str, Any] | None) -> bool:
    return bool(flags) or any(not c.get("ok", True) for c in (extraction or {}).get("checks") or [])


def diff(old: list[dict[str, Any]], new: list[dict[str, Any]]) -> dict[str, int]:
    """How much changed between two versions – for the audit log."""
    before = {r["id"]: r for r in old}
    after_ids = {r["id"] for r in new}
    cells = 0
    for row in new:
        prev = before.get(row["id"])
        if prev is None:
            continue
        a, b = prev.get("cells") or {}, row.get("cells") or {}
        for key in set(a) | set(b):
            if (a.get(key) or {}).get("f", (a.get(key) or {}).get("v")) != (b.get(key) or {}).get("f", (b.get(key) or {}).get("v")):
                cells += 1
    return {
        "cells": cells,
        "rows_added": sum(1 for r in new if r["id"] not in before),
        "rows_deleted": sum(1 for rid in before if rid not in after_ids),
    }
